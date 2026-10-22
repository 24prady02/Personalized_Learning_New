"""
run_answer_graph_100.py — NLI graph matching (src/lp_percpt/answer_graph.py, mode "nli":
NLI has_level edges + code-run edges, no LLM judge) on the 100 answers of
output/selected_100/classify_100_classified.xlsx, matched against the levels there.

Usage:
    python scripts/run_answer_graph_100.py [--nli cross-encoder/nli-deberta-v3-small] [--tag small]

Writes output/answer_graph_100_<tag>/results.json and accuracy.xlsx, and prints the match.
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
os.environ.pop("LP_JUDGE", None)       # NLI only: the LLM judge stays off

import openpyxl  # noqa: E402

from run_answer_graph_44 import run_and_question_edges  # noqa: E402
from run_student_scenarios import asked_part, load_scenarios  # noqa: E402
from src.lp_percpt.answer_graph import build  # noqa: E402
from src.lp_percpt.engine import LPPercptEngine  # noqa: E402

CLASSIFIED = ROOT / "output" / "selected_100" / "classify_100_classified.xlsx"
LV = ["L0", "L1", "L2", "L3", "L4", "none"]


def load_labels():
    ws = openpyxl.load_workbook(CLASSIFIED, read_only=True)["Classify 100"]
    return {r[0]: (str(r[5]).strip(), r[6]) for r in ws.iter_rows(values_only=True)
            if r[0] and str(r[0]).startswith("S") and len(str(r[0])) < 6}


def score(rows, mode):
    ref = [r["label"] for r in rows]
    got = [r[mode]["best"] or "none" for r in rows]
    n = len(rows)
    exact = sum(a == b for a, b in zip(ref, got))
    pairs = [(int(a[1]), int(b[1])) for a, b in zip(ref, got) if "none" not in (a, b)]
    within = sum(abs(a - b) <= 1 for a, b in pairs)
    mae = sum(abs(a - b) for a, b in pairs) / max(len(pairs), 1)
    try:
        from sklearn.metrics import cohen_kappa_score
        kappa = cohen_kappa_score(ref, got)
        qwk = cohen_kappa_score([a for a, _ in pairs], [b for _, b in pairs], weights="quadratic")
    except Exception:
        kappa = qwk = float("nan")
    per, hit, out = Counter(ref), Counter(a for a, b in zip(ref, got) if a == b), Counter(got)
    summary = [["Answers", n], ["Exact", f"{exact}/{n} ({exact / n:.0%})"],
               ["Within one level (both levelled)", f"{within}/{len(pairs)} ({within / max(len(pairs), 1):.0%})"],
               ["MAE (both levelled)", round(mae, 2)], ["Cohen kappa", round(kappa, 3)],
               ["Quadratic weighted kappa", round(qwk, 3)]]
    for x in LV:
        summary.append([f"{x} recall / precision",
                        f"{hit[x]}/{per[x]} ({hit[x] / per[x]:.0%})" if per[x] else "0/0",
                        f"{hit[x]}/{out[x]} ({hit[x] / out[x]:.0%})" if out[x] else "0/0"])
    conf = Counter(zip(ref, got))
    return summary, conf


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--nli", default="cross-encoder/nli-deberta-v3-small")
    ap.add_argument("--tag", default="small")
    args = ap.parse_args()
    out = ROOT / "output" / f"answer_graph_100_{args.tag}"
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    labels = load_labels()
    eng = LPPercptEngine(store_dir=out / "students", map_dir=None, review_log=None)
    kg = eng.kg
    questions, scenarios = load_scenarios()
    by_id = {s["id"]: s for s in scenarios}
    rows = []
    for sid, (label, why) in labels.items():
        s = by_id.get(sid)
        if s is None:
            print(sid, "not in the scenario workbook, skipped")
            continue
        q = questions.get(s["question"], {})
        asked = asked_part(q.get("question", ""), s["scenario_type"])
        shown = f"{q.get('source') or ''}\n{q.get('question', '')}\n{q.get('code', '')}"
        r = eng.process(sid, s["text"], code=q.get("code", ""), shown_text=shown, session_id="ag", asked=asked)
        run, ques = run_and_question_edges(r)
        row = {"id": sid, "label": label, "label_why": why, "text": s["text"], "asked": asked}
        for mode, qe in (("nli", []), ("full", ques)):
            g = build(s["text"], asked, q.get("code", ""), kg, args.nli, run_edges=run, question_edges=qe)
            main = g.main_concept()
            row[mode] = {"levels": g.levels, "reasons": g.reasons, "edges": [e.to_dict() for e in g.edges],
                         "best": g.best_level(), "main": main, "main_level": g.levels.get(main) if main else None}
        rows.append(row)
        print(sid, label, "| nli", row["nli"]["best"], "| full", row["full"]["best"], flush=True)

    (out / "results.json").write_text(json.dumps(rows, indent=1, ensure_ascii=False), encoding="utf-8")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for mode in ("nli", "full"):
        summary, conf = score(rows, mode)
        print(f"\n== {mode} ==")
        for line in summary:
            print("  " + "  ".join(str(x) for x in line))
        print("  confusion (rows = classified label, cols = graph matching):")
        print("        " + " ".join(f"{x:>5s}" for x in LV))
        for a in LV:
            print(f"  {a:>5s} " + " ".join(f"{conf[(a, b)]:5d}" for b in LV))
        ws = wb.create_sheet(f"Summary {mode}")
        for line in summary:
            ws.append(line)
        ws.append([])
        ws.append(["label \\ graph"] + LV)
        for a in LV:
            ws.append([a] + [conf[(a, b)] for b in LV])
    ws = wb.create_sheet("Per answer")
    ws.append(["ID", "Student answer", "Classified level", "NLI graph level", "Match", "Full mode level",
               "Main concept", "Graph reason (best edge)", "Classified why"])
    for r in rows:
        m = r["nli"]
        top = next((c for c, l in m["levels"].items() if l == m["best"]), None)
        ws.append([r["id"], r["text"], r["label"], m["best"] or "none", (m["best"] or "none") == r["label"],
                   r["full"]["best"] or "none", m["main"], m["reasons"].get(top, "") if top else "", r["label_why"]])
    wb.save(out / "accuracy.xlsx")
    print("\nwrote", out)


if __name__ == "__main__":
    main()
