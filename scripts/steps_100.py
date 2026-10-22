"""
steps_100.py — the output of every step for each of the 100 answers, from a run
of the classifier with the LLM judge (default: scenario_run_100_judge_no_nli,
NLI off), next to the levels in output/selected_100/classify_100_classified.xlsx.

  Step 1  semantic matching   the concepts linked to the answer (link score, how linked)
  Step 2  code run            what running the code confirmed or contradicted
  Step 3  LLM judge           level, way, quote, correct, why for each concept
  Step 4  fusion              the concept's level and which step set it
  Step 5  relation edges      (LP_JUDGE_RELATIONS runs) each prerequisite/parallel edge between the
                              concepts: does the student state it (Claude + quote found in the answer)
  Answer  final level         highest concept level, matched against the label

Usage: python scripts/steps_100.py [run folder]
Writes <run>/steps_100.xlsx.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[1]
sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def main() -> None:
    run = ROOT / "output" / (sys.argv[1] if len(sys.argv) > 1 else "scenario_run_100_judge_no_nli")
    ws = openpyxl.load_workbook(ROOT / "output" / "selected_100" / "classify_100_classified.xlsx",
                                read_only=True)["Classify 100"]
    ref = {r[0]: (str(r[5]).strip(), r[6]) for r in ws.iter_rows(values_only=True)
           if r[0] and str(r[0]).startswith("S") and len(str(r[0])) < 6}
    res = {s["id"]: s for s in json.loads((run / "results.json").read_text(encoding="utf-8"))["scenarios"]}

    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "Per answer"
    a.append(["ID", "Student answer", "Classified level", "Final level", "Match", "Main concept",
              "Set by", "Concepts linked (step 1)", "Code-run facts (step 2)", "Judge levels (step 3)",
              "Classified why", "Gates", "Student-stated edges (step 5)"])
    c_ws = wb.create_sheet("Per concept (steps)")
    c_ws.append(["ID", "Concept", "Step 1 link score", "Step 1 linked by", "Step 2 code-run facts",
                 "Step 3 judge level", "Step 3 judge way", "Step 3 way id", "Stored in the map", "Step 3 quote", "Step 3 correct", "Step 3 why",
                 "Step 4 level", "Step 4 set by", "Step 4 reason"])
    r_ws = wb.create_sheet("Relation edges (step 5)")
    r_ws.append(["ID", "Source", "Relation", "Target", "Claude: stated", "Quote", "Quote found in answer",
                 "Student edge (stated + quote found + both ends stored)"])
    for sid, (label, why) in ref.items():
        s = res.get(sid)
        if s is None:
            continue
        final = s["max_level"] or "none"
        main, linked, judged, facts_all = None, [], [], []
        for c in s["concepts"]:
            j = c.get("judge") if isinstance(c.get("judge"), dict) else {}
            facts = [f"{e['statement']}" for e in c.get("evidence", [])
                     if e["statement"].startswith("confirmed by the run") or e["check"] in ("wrong_trace", "own_run")
                     or (e["check"] == "held" and "match a run" in e["statement"])]
            facts_all += facts
            how = "judge only (closest rules)" if "judge_only" in c.get("gates", []) else "concept linker"
            linked.append(f"{c['concept']} ({c.get('link_score', 0):.2f})")
            if j:
                judged.append(f"{c['concept']}: {j.get('level')}")
            if main is None and c.get("prompt_level") == s["max_level"] and c.get("level_source"):
                main = c
            c_ws.append([sid, c["concept"], c.get("link_score"), how, "; ".join(facts),
                         j.get("level"), j.get("way"), j.get("way_id"),
                         "no: " + ", ".join(c.get("gates") or []) if c.get("gates") else "yes", j.get("quote"), j.get("correct"), j.get("why"),
                         c.get("prompt_level"), c.get("level_source") or "", c.get("prompt_reason") or c.get("stop_reason")])
        a.append([sid, s["text"], label, final, final == label, (main or {}).get("concept"),
                  (main or {}).get("level_source"), ", ".join(linked), "; ".join(dict.fromkeys(facts_all)),
                  ", ".join(judged), why, ", ".join(s.get("prompt_gates") or []),
                  "; ".join(f"({e['source']}, {e['type']}, {e['target']})" for e in s.get("student_edges") or []
                            if e["student_edge"])])
        for e in s.get("student_edges") or []:
            r_ws.append([sid, e["source"], e["type"], e["target"], e["stated"], e["quote"], e["quote_found"],
                         e["student_edge"]])
    for sh in (a, c_ws, r_ws):
        for cell in sh[1]:
            cell.font = Font(bold=True)
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col, width in zip("ABCDEFGHIJKLMNO", (7, 50, 10, 10, 8, 20, 12, 40, 50, 30, 50, 14, 50, 14, 50)):
            sh.column_dimensions[col].width = width
        sh.freeze_panes = "A2"
    out = run / "steps_100.xlsx"
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
