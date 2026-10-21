"""
trace_steps_44.py — for the 44 labelled scenarios, record what each step of the
classifier produces, so each step can be checked on its own:

  1. Semantic matching (MiniLM): per clause, the closest concepts by meaning and
     the closest misconception, before any cue or threshold.
  2. Concept linking: the score of each linked concept, split into its parts
     (code cue, word cue, asked part, meaning, misconception).
  3. NLI (deberta-v3-small) on each linked concept: best fact entailment, best
     misconception entailment and best way rule (L0 / L3 / L4) per clause, and
     whether it passed the threshold.
  4. Evidence and level: checks E1-E4, the way that set the level, expected level.

Writes output/scenario_run_v7/steps_44.json and steps_44.xlsx.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import numpy as np  # noqa: E402
import openpyxl  # noqa: E402

from scripts.run_student_scenarios import asked_part, load_scenarios  # noqa: E402
from src.lp_percpt import models  # noqa: E402
from src.lp_percpt.evidence import _desc, _score, _way_hit  # noqa: E402
from src.lp_percpt.identify import _cue_hits, identify, split_clauses  # noqa: E402
try:
    from src.lp_percpt.identify import question_concepts  # noqa: E402
except ImportError:  # v7 linking has no question prior
    question_concepts = lambda asked, code, kg: {}  # noqa: E731
from src.lp_percpt.knowledge import get_knowledge  # noqa: E402
from src.lp_percpt.thresholds import (FACT_ENT, HELD_ENT, HELD_SIM, MIS_FLOOR, SEM_FLOOR,  # noqa: E402
                                      W_CODE, W_CODE_SHOWN, W_MIS, W_SEM, W_WORD, WAY_ENT, WAY_SIM)
W_QUESTION = 0

OUT = ROOT / "output" / (sys.argv[1] if len(sys.argv) > 1 else "scenario_run_v7")


def main() -> None:
    kg = get_knowledge()
    questions, scenarios = load_scenarios()
    v7 = {s["id"]: s for s in json.loads((OUT / "results.json").read_text(encoding="utf-8"))["scenarios"]}
    labelled = [s for s in scenarios if s["expected_level"]]

    concepts = list(kg.concepts.values())
    m_targets, m_owner = [], []
    for c in concepts:
        for s in [f"{c.label.replace('_', ' ')}: {c.description}"] + c.facts:
            m_targets.append(s); m_owner.append(c.id)
    mis_ids = list(kg.misconceptions)
    mis_texts = [kg.misconceptions[m].wrong_belief for m in mis_ids]

    rows = []
    for s in labelled:
        q = questions.get(s["question"], {})
        code, asked = q.get("code", ""), asked_part(q.get("question", ""), s["scenario_type"])
        text = s["text"].strip()
        clauses = split_clauses(text)

        # 1 — semantic matching, per clause
        sem_steps = []
        if clauses:
            sim = models.similarity(clauses, m_targets)
            msim = models.similarity(clauses, mis_texts)
            for i, cl in enumerate(clauses):
                best = {}
                for j, owner in enumerate(m_owner):
                    if sim[i, j] > best.get(owner, (0, ""))[0]:
                        best[owner] = (float(sim[i, j]), m_targets[j])
                top = sorted(best.items(), key=lambda kv: -kv[1][0])[:3]
                k = int(np.argmax(msim[i]))
                sem_steps.append({"clause": cl,
                                  "top_meaning": [(cid, round(v, 3), t[:90]) for cid, (v, t) in top],
                                  "top_misconception": (mis_ids[k], round(float(msim[i, k]), 3))})

        # 2 — concept linking with the score split into its parts
        hits = identify(text, kg, code=code, clauses=clauses, asked=asked)
        qc = question_concepts(asked, code, kg)
        link = []
        for h in hits:
            c = kg.concepts[h.concept]
            code_h, shown_h, word_h = _cue_hits(c, text.lower(), text, code)
            link.append({"concept": h.concept, "score": h.score,
                         "code": W_CODE * min(len(code_h), 2), "shown_code": W_CODE_SHOWN * min(len(shown_h), 1),
                         "words": W_WORD * sum(2 if " " in w.strip() else 1 for w in word_h[:3]),
                         "asked": round(W_QUESTION * getattr(h, "question_prior", 0.0), 3),
                         "meaning": round(W_SEM * max(0.0, h.meaning_sim - SEM_FLOOR), 3),
                         "misconception": round(W_MIS * max(0.0, h.misconception_sim - MIS_FLOOR), 3),
                         "meaning_sim": h.meaning_sim, "cues": code_h + shown_h + word_h})

        # 3 — NLI on each linked concept
        nli_steps = []
        for h in hits:
            c = kg.concepts[h.concept]
            if not clauses:
                break
            sc = _score(c, kg, clauses)
            mis = sorted(c.misconceptions)
            for i, cl in enumerate(clauses):
                f = max(((r["entailment"], r["contradiction"], c.facts[j]) for (ii, j), r in sc.nli_f.items() if ii == i),
                        default=(0.0, 0.0, ""))
                m = max(((r["entailment"], float(sc.mis[i, j]), mis[j]) for (ii, j), r in sc.nli_m.items() if ii == i),
                        default=(0.0, 0.0, ""))
                ways = {}
                for lv in ("L0", "L3", "L4"):
                    cand = [(r["entailment"], float(sc.way[lv][i, j]), sc.way_of[lv][j][0])
                            for (k, ii, j), r in sc.nli_w.items() if k == lv and ii == i]
                    if cand:
                        e, sm, rule = max(cand)
                        hit = _way_hit(sc, lv, i, cl, kg)
                        ways[lv] = {"ent": round(e, 3), "sim": round(sm, 3), "rule": rule[:110],
                                    "passed": hit is not None}
                nli_steps.append({"concept": c.id, "clause": cl,
                                  "fact_ent": round(f[0], 3), "fact_contra": round(f[1], 3), "fact": f[2][:110],
                                  "fact_passed": f[0] >= FACT_ENT,
                                  "mis_ent": round(m[0], 3), "mis_sim": round(m[1], 3), "mis": m[2],
                                  "mis_passed": m[0] >= HELD_ENT and m[1] >= HELD_SIM,
                                  "ways": ways})

        r = v7[s["id"]]
        assessed = [c for c in r["concepts"] if c["assessed"]]
        rows.append({"id": s["id"], "type": s["scenario_type"], "expected": s["expected_level"],
                     "text": text, "asked": asked, "question_concepts": qc, "clauses": clauses,
                     "semantic": sem_steps, "linking": link, "nli": nli_steps,
                     "final": [{"concept": c["concept"], "level": c["prompt_level"], "checks": c["checks"],
                                "reason": c["prompt_reason"], "gates": c["gates"], "held": c["held"]}
                               for c in assessed],
                     "not_assessed": [c["concept"] for c in r["concepts"] if not c["assessed"]],
                     "prompt_gates": r["prompt_gates"], "top_concept": r["top_concept"], "top_level": r["top_level"]})
        print(s["id"], "done", flush=True)

    (OUT / "steps_44.json").write_text(json.dumps(rows, indent=1, ensure_ascii=False), encoding="utf-8")
    write_xlsx(rows)
    print("wrote", OUT / "steps_44.json", OUT / "steps_44.xlsx")


def write_xlsx(rows) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "1 Semantic matching"
    ws.append(["Scenario", "Expected", "Clause", "Meaning #1", "sim", "Meaning #2", "sim", "Meaning #3", "sim",
               "Closest misconception", "sim"])
    for r in rows:
        for st in r["semantic"]:
            t = st["top_meaning"] + [("", "", "")] * (3 - len(st["top_meaning"]))
            ws.append([r["id"], r["expected"], st["clause"], t[0][0], t[0][1], t[1][0], t[1][1], t[2][0], t[2][1],
                       *st["top_misconception"]])
    ws = wb.create_sheet("2 Concept linking")
    ws.append(["Scenario", "Expected", "Asked part", "Concept", "Score", "Code cue", "Shown code", "Word cue",
               "Question prior", "Meaning", "Misconception", "Cues"])
    for r in rows:
        for l in r["linking"] or [{"concept": "(none linked)"}]:
            ws.append([r["id"], r["expected"], r["asked"][:120], l["concept"], l.get("score"), l.get("code"),
                       l.get("shown_code"), l.get("words"), l.get("asked"), l.get("meaning"), l.get("misconception"),
                       ", ".join(l.get("cues", []))[:120]])
    ws = wb.create_sheet("3 NLI")
    ws.append(["Scenario", "Concept", "Clause", "Fact ent", "Fact contra", "Fact passed", "Best fact",
               "Mis ent", "Mis sim", "Mis passed", "Misconception",
               "L0 ent", "L0 passed", "L3 ent", "L3 passed", "L3 rule", "L4 ent", "L4 passed", "L4 rule"])
    for r in rows:
        for n in r["nli"]:
            w = n["ways"]
            g = lambda lv, k: w.get(lv, {}).get(k, "")
            ws.append([r["id"], n["concept"], n["clause"][:150], n["fact_ent"], n["fact_contra"], n["fact_passed"],
                       n["fact"], n["mis_ent"], n["mis_sim"], n["mis_passed"], n["mis"],
                       g("L0", "ent"), g("L0", "passed"), g("L3", "ent"), g("L3", "passed"), g("L3", "rule"),
                       g("L4", "ent"), g("L4", "passed"), g("L4", "rule")])
    ws = wb.create_sheet("4 Level")
    ws.append(["Scenario", "Type", "Expected", "Got", "Match", "Top concept", "Why this level", "E1", "E2", "E3", "E4",
               "Not assessed", "Prompt gates"])
    for r in rows:
        f = r["final"][0] if r["final"] else {}
        ch = f.get("checks", {})
        ws.append([r["id"], r["type"], r["expected"], r["top_level"], r["top_level"] == r["expected"], r["top_concept"],
                   f.get("reason", ""), ch.get("E1"), ch.get("E2"), ch.get("E3"), ch.get("E4"),
                   ", ".join(r["not_assessed"]), ", ".join(r["prompt_gates"])])
    wb.save(OUT / "steps_44.xlsx")


if __name__ == "__main__":
    main()
