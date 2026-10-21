"""
export_100_for_llm.py — the 100 selected scenarios (output/selected_100) as one
Markdown document to give an LLM, which classifies every student answer at L0-L4 / none
from the graph's level definitions and boundary rules.

The document never shows a label, the algorithm's level, or the scenario type (the type
names the level, e.g. "J7 L0 way"). The labels go to a separate answer key for scoring.

Writes output/<set>/classify_100.md, classify_100.xlsx (one sheet) and answer_key_100.json
Usage: python scripts/export_100_for_llm.py [--set selected_100]
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from reference_labels_365 import BOUNDARIES, level_definitions  # noqa: E402
from run_student_scenarios import asked_part, load_scenarios  # noqa: E402


INSTRUCTIONS = """# Classifying 100 student answers on the Chapter-1 learning ladder (L0-L4)

You are an experienced Java teacher. For each of the 100 student answers below, give the level the answer
shows on a 5-level learning ladder (L0-L4), or "none" if the answer makes no claim.

How to decide:
- Use only the level definitions and boundary rules below.
- Work out what the code really does (its run output is given where it could be run) and whether the
  student's reasoning is correct before you decide.
- Mark the answer as a whole, for the main thing the question part asks.
- Judge each answer on its own.

Reply with one table and nothing else, one row per answer, in the order given:

| ID | Level | Why (one sentence) |
|----|-------|--------------------|
| S088 | L2 | ... |
"""


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--set", default="selected_100", help="folder under output/ holding selected_60.json")
    OUT = ROOT / "output" / ap.parse_args().set
    chosen = json.loads((OUT / "selected_60.json").read_text(encoding="utf-8"))
    questions, scenarios = load_scenarios()
    scen = {s["id"]: s for s in scenarios}
    res = {s["id"]: s for s in json.loads((ROOT / "output" / "scenario_run_v10_all" / "results.json")
                                         .read_text(encoding="utf-8"))["scenarios"]}
    parts = [INSTRUCTIONS, "## Level definitions\n\n" + level_definitions(), "## " + BOUNDARIES, "## The 100 answers"]
    key, rows = [], []
    for n, c in enumerate(sorted(chosen, key=lambda c: int(c["id"][1:])), 1):
        s = scen[c["id"]]
        q = questions.get(s["question"], {})
        asked = asked_part(q.get("question", ""), s["scenario_type"]) or q.get("question", "")
        run = res[c["id"]].get("run_output")
        run = str(run).strip() if run not in (None, "None", "") else "(not run: no runnable code, or it needs input)"
        if len(run) > 600:                             # as reference_labels_365.material; an endless loop prints without end
            run = run[:600] + f"\n... (output cut: {len(run)} characters in all)"
        code = (q.get("code") or "").strip() or "(none)"
        rows.append([c["id"], asked.strip(), code, run, s["text"].strip(), "", ""])
        parts.append(f"### {n}. {c['id']}\n\n**Question (the part the student answers):**\n\n{asked.strip()}\n\n"
                     f"**Code:**\n\n```java\n{code}\n```\n\n**Output of running the code:**\n\n```\n{str(run).strip()}\n```\n\n"
                     f"**Student answer:**\n\n> " + s["text"].strip().replace("\n", "\n> "))
        key.append({"id": c["id"], "level": c["level"], "level_by": "workbook (LLM)" if c["human"] else "Fable",
                    "algorithm": c["algorithm"]})
    (OUT / "classify_100.md").write_text("\n\n".join(parts) + "\n", encoding="utf-8")
    (OUT / "answer_key_100.json").write_text(json.dumps(key, indent=1, ensure_ascii=False), encoding="utf-8")
    write_sheet(OUT / "classify_100.xlsx", "\n\n".join(parts[:3]), rows)
    print(f"wrote {OUT / 'classify_100.md'}, classify_100.xlsx ({len(key)} answers) and answer_key_100.json")


def write_sheet(path: Path, instructions: str, rows) -> None:
    """One sheet: the instructions and level rules on top, then one row per answer with
    empty Level / Why columns for the LLM to fill."""
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Classify 100"
    text = instructions.replace("Reply with one table and nothing else, one row per answer, in the order given:",
                                "Fill the Level and Why columns for every row (or reply with a table: ID | Level | Why).")
    text = text.split("| ID | Level |")[0].rstrip() + "\n\n" + instructions.split("## Level definitions", 1)[1].join(["## Level definitions", ""])
    for line in text.splitlines():
        ws.append([line.lstrip("#").strip()])
        if line.startswith("#"):
            ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    ws.append(["ID", "Question (the part the student answers)", "Code", "Output of running the code",
               "Student answer", "Level (L0-L4 / none)", "Why (one sentence)"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    head = ws.max_row
    for r in rows:
        ws.append(r)
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFG", (8, 45, 50, 35, 50, 14, 45)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(head + 1, 1)
    wb.save(path)


if __name__ == "__main__":
    main()
