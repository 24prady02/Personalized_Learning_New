"""
run_student_scenarios.py — run docs/Chapter1_student_scenarios_set.xlsx through
the LP classifier (src/lp_percpt) and record every classification.

Each scenario is one prompt: the student's reasoning is the prompt text and the
question's code is the code it reasons about (used to link concepts). The
question and code are also passed as `shown_text`, so text or code copied from
them is not counted as evidence (parroting and pasted-code gates).
Each scenario is its own student, except rows whose Note names the same student
("Student S1"), which share one student graph. Every student gets a concept map.

Outputs (output/scenario_run/):
    results.json           everything the engine returned, per scenario
    results.xlsx           Scenarios / Concepts / Scenario types / Level check
    maps/<student>.html    the concept map of each student

Usage:
    python scripts/run_student_scenarios.py [--limit N]
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import openpyxl  # noqa: E402

from src.lp_percpt.engine import LPPercptEngine  # noqa: E402
from src.lp_percpt.execution import get_tracer  # noqa: E402

WORKBOOK = ROOT / "docs" / "Chapter1_student_scenarios_set.xlsx"
OUT = ROOT / "output" / "scenario_run"
_STUDENT_RE = re.compile(r"\bStudent (S\d+)\b")
_LEVEL_RE = re.compile(r"\bL([0-4])\b")
_PART_RE = re.compile(r"\bPart \(([a-h])\)")
_PARTS_RE = re.compile(r"(?:^|\s)\(([a-h])\)\s")


def asked_part(question: str, scenario_type: str) -> str:
    """The question part the student answers: the part the scenario type names
    ("J4 Part (d) L4" -> the text of (d)); the whole question when it has no
    parts; nothing when it has parts but the scenario does not say which."""
    marks = list(_PARTS_RE.finditer(question or ""))
    if not marks:
        return question or ""
    m = _PART_RE.search(scenario_type or "")
    if not m:
        return ""
    for k, mk in enumerate(marks):
        if mk.group(1) == m.group(1):
            end = marks[k + 1].start() if k + 1 < len(marks) else len(question)
            return question[mk.end():end].strip()
    return ""


def _unquote(text: str) -> str:
    t = (text or "").strip()
    if len(t) >= 2 and t[0] == t[-1] == '"':
        t = t[1:-1]
    return t.replace("\\n", "\n")


def load_scenarios():
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True)
    questions = {}
    for row in list(wb["Questions"].iter_rows(values_only=True))[1:]:
        qid, fmt, source, code, ask = row[:5]
        questions[qid] = {"format": fmt, "source": source, "code": code or "", "question": ask or ""}
    scenarios = []
    for row in list(wb["Student scenarios"].iter_rows(values_only=True))[1:]:
        sid, qid, origin, stype, kind, writes, note = row[:7]
        if not sid:
            continue
        m = _STUDENT_RE.search(note or "")
        lv = _LEVEL_RE.search(stype or "")
        scenarios.append({"id": sid, "question": qid, "from": origin, "scenario_type": stype,
                          "kind": kind, "text": _unquote(writes), "note": note or "",
                          "student": m.group(1) if m else sid,
                          "expected_level": f"L{lv.group(1)}" if lv else None})
    return questions, scenarios


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--limit", type=int, help="run only the first N scenarios")
    ap.add_argument("--labelled", action="store_true", help="run only the scenarios with an expected level")
    ap.add_argument("--ids", help="JSON file with a list of scenario ids to run")
    ap.add_argument("--out", help="output folder (default output/scenario_run)")
    ap.add_argument("--sheets-only", action="store_true",
                    help="rebuild results.xlsx from an earlier run's results.json and student graphs")
    args = ap.parse_args()
    global OUT
    if args.out:
        OUT = Path(args.out).resolve()
    if args.sheets_only:
        from src.lp_percpt.knowledge import get_knowledge
        results = json.loads((OUT / "results.json").read_text(encoding="utf-8"))["scenarios"]
        write_xlsx(results, get_knowledge())
        print(f"wrote {OUT / 'results.xlsx'}")
        return
    questions, scenarios = load_scenarios()
    if args.labelled:
        scenarios = [s for s in scenarios if s["expected_level"]]
    if args.ids:
        keep = set(json.loads(Path(args.ids).read_text(encoding="utf-8")))
        scenarios = [s for s in scenarios if s["id"] in keep]
    if args.limit:
        scenarios = scenarios[:args.limit]
    if OUT.exists():
        shutil.rmtree(OUT)
    eng = LPPercptEngine(store_dir=OUT / "students", map_dir=OUT / "maps", review_log=None)
    kg = eng.kg

    def shown_text(q):
        # the question as the student saw it: source line (it may give values such as
        # "value = 2"), the question, and the code
        return f"{q.get('source') or ''}\n{q.get('question', '')}\n{q.get('code', '')}"

    # each question's code (and its misconception runs) is run the first time one of
    # its scenarios comes up, then reused from the cache
    tracer = get_tracer()

    def student_edges(relations, assessed, text):
        # LP_JUDGE_RELATIONS=1: a graph edge is the student's own when the judge says the answer
        # states it, both ends are stored concepts, and the quote is found in the answer
        norm = lambda t: " ".join(re.findall(r"\w+", (t or "").lower()))
        stored = {c["concept"] for c in assessed if "judge_no_level" not in (c.get("gates") or [])}
        out = []
        for rel in relations or []:
            m = re.match(r"\s*(\w+)\s*-(\w+)->\s*(\w+)", str(rel.get("edge", "")))
            if not m:
                continue
            a, typ, b = m.groups()
            quoted = bool(norm(rel.get("quote"))) and norm(rel.get("quote")) in norm(text)
            out.append({"source": a, "type": typ, "target": b, "quote": rel.get("quote", ""),
                        "stated": bool(rel.get("stated")), "quote_found": quoted,
                        "student_edge": bool(rel.get("stated")) and quoted and a in stored and b in stored})
        return out

    results = []
    for i, s in enumerate(scenarios, 1):
        q = questions.get(s["question"], {})
        # prompt = the student's reasoning; code = the question's code it reasons about.
        # The code is also shown text, so copying it back earns nothing (G7, G8).
        # the part asked is known to the tutor; the level in the scenario type is not passed
        r = eng.process(s["student"], s["text"], code=q.get("code", ""), shown_text=shown_text(q),
                        session_id="scenarios", asked=asked_part(q.get("question", ""), s["scenario_type"]))
        concepts = [{k: c.get(k) for k in ("concept", "label", "chapter", "link_score", "prompt_level", "level",
                                           "provisional", "confidence", "gates", "checks", "held", "refuted", "assessed",
                                           "stop_reason", "prompt_reason", "follow_up", "evidence",
                                           "level_source", "wrong_rule", "judge")} for c in r["concepts"]]
        # levels come from assessed concepts only: a linked concept the prompt shows
        # nothing about has prompt_level L0 meaning "not assessed", not prestructural
        assessed = [c for c in concepts if c["assessed"]]
        top = assessed[0] if assessed else None
        trace = tracer.trace(q.get("code", ""), hint=shown_text(q)) if tracer.available and q.get("code") else None
        max_level = ("L0" if any(c.get("wrong_rule") for c in assessed)
                     else max((c["prompt_level"] for c in assessed), default=None))
        jd = r.get("judge") or {}
        # rubric judge: its one level for the answer as a whole, unless the code run
        # proves a wrong trace or a reproduced misconception (L0) or the prompt makes no claim
        run_l0 = any(c.get("level_source") == "run" and c["prompt_level"] == "L0" for c in assessed)
        if r["prompt_gates"] and set(r["prompt_gates"]) & {"no_claim", "help_seeking"}:
            final = None
        elif os.environ.get("LP_FINAL") == "max_level" and max_level:
            # LP_FINAL=max_level: the highest level among the stored concepts (L0 if one holds a wrong rule)
            final = max_level
        elif run_l0:
            final = "L0"
        elif jd.get("answer_level") in ("L0", "L1", "L2", "L3", "L4"):
            final = jd["answer_level"]
        elif jd.get("answer_level") == "none":
            final = None
        else:
            final = max_level
        results.append({**s, "question_text": q.get("question", ""), "code": q.get("code", ""),
                        "item_format": q.get("format", ""), "prompt_gates": r["prompt_gates"],
                        "judge_unavailable": jd.get("unavailable"),
                        "judge_answer": {k: jd.get(k) for k in ("answer_level", "main_concept", "why", "concepts", "relations")},
                        "final_level": final,
                        "run_output": trace.run.output if trace else None,
                        "concepts": concepts,
                        "top_concept": top["concept"] if top else None,
                        "top_level": top["prompt_level"] if top else None,
                        # an answer holding a rule that contradicts Java is L0 (the graph's L0 edge
                        # logic), when the code run or the judge found it
                        "max_level": max_level,
                        "student_edges": student_edges(jd.get("relations"), assessed, s["text"]),
                        "map": f"maps/{s['student']}.html"})
        print(f"[{i}/{len(scenarios)}] {s['id']} {s['question']}: "
              + (", ".join(f"{c['concept']}={c['prompt_level']}" for c in concepts) or "(" + ",".join(r["prompt_gates"]) + ")"), flush=True)

    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "results.json").write_text(json.dumps({"graph_version": kg.version, "scenarios": results},
                                                 indent=1, ensure_ascii=False), encoding="utf-8")
    write_xlsx(results, kg)
    print(f"\nwrote {OUT / 'results.json'}, {OUT / 'results.xlsx'} and {len(set(s['student'] for s in scenarios))} concept maps")


def _safe(student_id: str) -> str:
    """The file name the engine uses for a student graph."""
    return "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in student_id)[:80] or "anon"


def write_xlsx(results, kg) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scenarios"
    ws.append(["ID", "Question", "Scenario type", "Student writes", "Prompt gates", "Concepts (prompt level)",
               "Top concept", "Top level", "Highest level", "Expected level (from type)", "Held misconceptions",
               "Gates per concept"])
    for r in results:
        ws.append([r["id"], r["question"], r["scenario_type"], r["text"], ", ".join(r["prompt_gates"]),
                   "; ".join(f"{c['concept']} {c['prompt_level']}" for c in r["concepts"]),
                   r["top_concept"], r["top_level"], r["max_level"], r["expected_level"],
                   ", ".join(sorted({m for c in r["concepts"] for m in c["held"]})),
                   "; ".join(f"{c['concept']}: {','.join(c['gates'])}" for c in r["concepts"] if c["gates"])])

    ws = wb.create_sheet("Concept nodes")
    ws.append(["Scenario", "Student", "Question", "Concept node", "Section", "Status on the map",
               "Level on the map", "Level from this prompt", "Held misconceptions", "At risk from",
               "Why this level", "Follow-up question", "Evidence"])
    for r in results:
        g_path = OUT / "students" / f"{_safe(r['student'])}.json"
        nodes = json.loads(g_path.read_text(encoding="utf-8"))["nodes"] if g_path.exists() else {}
        in_prompt = {c["concept"]: c for c in r["concepts"]}
        order = list(in_prompt) + sorted(k for k in nodes if k not in in_prompt)
        for cid in order:
            n, c = nodes.get(cid, {}), in_prompt.get(cid)
            if c is not None:
                status = "assessed in this prompt" if c.get("assessed", True) else "linked, not assessed"
            elif n.get("observed"):
                status = "assessed in an earlier prompt"
            elif n.get("inferred_level"):
                status = "inferred (prerequisite)"
            elif n.get("frontier"):
                status = "frontier (next to assess)"
            else:
                continue
            level = (f"L{n['level']}" if n.get("observed") else
                     f"L{n['inferred_level']} (inferred)" if n.get("inferred_level") else "")
            ws.append([r["id"], r["student"], r["question"], cid, kg.concepts[cid].chapter if cid in kg.concepts else "",
                       status, level,
                       (c["prompt_level"] if c is not None and c.get("assessed", True) else ""),
                       ", ".join(c["held"]) if c else ", ".join(m for m, s in n.get("misconceptions", {}).items() if s.get("status") == "held"),
                       ", ".join(n.get("at_risk", [])),
                       (c["prompt_reason"] if c is not None else n.get("stop_reason", "")),
                       (c.get("follow_up") or "") if c is not None else "",
                       "; ".join(f"{e['check']}: {e['statement'][:120]}" for e in (c["evidence"] if c else [])
                                 if e["check"] in ("E1", "E2", "E3", "E4", "L0", "held", "wrong_trace"))[:1000]])

    ws = wb.create_sheet("Concepts")
    ws.append(["Concept", "Section", "Times linked", "Not assessed", "L0", "L1", "L2", "L3", "L4",
               "Held misconceptions (count)"])
    per = defaultdict(Counter)
    held = defaultdict(Counter)
    for r in results:
        for c in r["concepts"]:
            per[c["concept"]][c["prompt_level"] if c.get("assessed", True) else "n/a"] += 1
            held[c["concept"]].update(c["held"])
    for cid, cnt in sorted(per.items(), key=lambda x: -sum(x[1].values())):
        ws.append([cid, kg.concepts[cid].chapter, sum(cnt.values()), cnt["n/a"]] + [cnt[f"L{k}"] for k in range(5)]
                  + ["; ".join(f"{m} ({n})" for m, n in held[cid].most_common())])

    ws = wb.create_sheet("Scenario types")
    ws.append(["Scenario type", "Scenarios", "Unlinked / gated", "L0", "L1", "L2", "L3", "L4"])
    by = defaultdict(list)
    for r in results:
        by[r["scenario_type"]].append(r)
    for t, rs in sorted(by.items(), key=lambda x: -len(x[1])):
        cnt = Counter(r["top_level"] for r in rs)
        ws.append([t, len(rs), cnt[None]] + [cnt[f"L{k}"] for k in range(5)])

    ws = wb.create_sheet("Level check")
    ws.append(["ID", "Scenario type", "Expected", "Top level", "Highest level", "Match (highest)", "Student writes"])
    for r in results:
        if r["expected_level"]:
            ws.append([r["id"], r["scenario_type"], r["expected_level"], r["top_level"], r["max_level"],
                       "yes" if r["max_level"] == r["expected_level"] else "no", r["text"]])
    wb.save(OUT / "results.xlsx")


if __name__ == "__main__":
    main()
