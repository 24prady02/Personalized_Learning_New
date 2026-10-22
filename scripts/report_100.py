"""
report_100.py — score a run on the 100 answers of output/selected_100/classify_100_classified.xlsx
for both parts of the matching:

  concept matching      linked concepts and the main concept vs the hand-checked concepts
                        (GOLD in check_steps_44.py, the answers that have one)
  level classification  the final level vs the classified level (exact, within one, per level,
                        confusion), and which step set it

and write every step's output per answer plus the concept-map data.

Usage: python scripts/report_100.py [run folder, default scenario_run_100_rubric]
Writes <run>/report_100.json and <run>/steps_100_rubric.xlsx; prints the scores.
"""
from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from check_steps_44 import GOLD  # noqa: E402
from src.lp_percpt.knowledge import get_knowledge  # noqa: E402

LV = ["L0", "L1", "L2", "L3", "L4", "none"]
GATES = {"no_claim", "help_seeking"}


def labels():
    ws = openpyxl.load_workbook(ROOT / "output" / "selected_100" / "classify_100_classified.xlsx",
                                read_only=True)["Classify 100"]
    return {r[0]: (str(r[5]).strip(), r[6]) for r in ws.iter_rows(values_only=True)
            if r[0] and str(r[0]).startswith("S") and len(str(r[0])) < 6}


def set_by(s):
    if set(s["prompt_gates"]) & GATES:
        return "no-claim check"
    if any(c.get("level_source") == "run" and c["prompt_level"] == "L0" for c in s["concepts"] if c["assessed"]):
        return "code run (L0)"
    a = (s.get("judge_answer") or {}).get("answer_level")
    if a in LV:
        return "LLM judge (overall)"
    return "per-concept fusion"


def run_facts(c):
    return [e["statement"] for e in c.get("evidence", [])
            if e["statement"].startswith("confirmed by the run") or e["check"] in ("wrong_trace", "own_run")
            or (e["check"] == "held" and "match a run" in e["statement"])]


def main() -> None:
    run = ROOT / "output" / (sys.argv[1] if len(sys.argv) > 1 else "scenario_run_100_rubric")
    ref = labels()
    kg = get_knowledge()
    res = {s["id"]: s for s in json.loads((run / "results.json").read_text(encoding="utf-8"))["scenarios"]}
    ids = [i for i in ref if i in res]

    rows, link_counts, shown_counts, pair_counts = [], Counter(), Counter(), Counter()
    for i in ids:
        s, (label, why) = res[i], ref[i]
        final = (s["final_level"] if "final_level" in s else s.get("max_level")) or "none"
        ja = s.get("judge_answer") or {}
        linked = [c for c in s["concepts"] if "judge_only" not in (c.get("gates") or [])]
        extra = [c for c in s["concepts"] if "judge_only" in (c.get("gates") or [])]
        jl = {c["concept"]: c["judge"].get("level") for c in s["concepts"] if isinstance(c.get("judge"), dict)}
        shown = [c for c, l in jl.items() if l in LV[:5]]
        top = max(linked, key=lambda c: c.get("link_score") or 0)["concept"] if linked else None
        for c in linked:
            link_counts[c["concept"]] += 1
        for c in shown:
            shown_counts[c] += 1
        names = sorted({c["concept"] for c in s["concepts"]})
        for a in names:
            for b in names:
                if a < b and (b in kg.concepts[a].neighbours or a in kg.concepts[b].neighbours):
                    pair_counts[(a, b)] += 1
        rows.append({
            "id": i, "text": s["text"], "label": label, "label_why": why, "final": final, "match": final == label,
            "set_by": set_by(s), "gates": s["prompt_gates"],
            "linked": [(c["concept"], round(c.get("link_score") or 0, 2)) for c in linked],
            "extra": [c["concept"] for c in extra], "top_linked": top,
            "run_facts": list(dict.fromkeys(f for c in s["concepts"] for f in run_facts(c))),
            "judge_concepts": jl, "judge_answer": ja.get("answer_level"), "judge_main": ja.get("main_concept"),
            "judge_why": ja.get("why"), "max_level": s.get("max_level") or "none",
            "gold": sorted(GOLD.get(i, [])),
        })

    # level classification
    n = len(rows)
    exact = sum(r["match"] for r in rows)
    pairs = [(int(r["label"][1]), int(r["final"][1])) for r in rows if "none" not in (r["label"], r["final"])]
    within = sum(abs(a - b) <= 1 for a, b in pairs)
    per, hit = Counter(r["label"] for r in rows), Counter(r["label"] for r in rows if r["match"])
    conf = Counter((r["label"], r["final"]) for r in rows)
    src, src_ok = Counter(r["set_by"] for r in rows), Counter(r["set_by"] for r in rows if r["match"])
    old_exact = sum(r["max_level"] == r["label"] for r in rows)
    # concept matching on the answers with hand-checked concepts
    g = [r for r in rows if r["gold"] and not set(r["gates"]) & GATES]
    in_map = sum(bool(set(c for c, _ in r["linked"]) & set(r["gold"])) for r in g)
    top_ok = sum(r["top_linked"] in r["gold"] for r in g)
    main_ok = sum(r["judge_main"] in r["gold"] for r in g)
    rep = {
        "n": n, "exact": exact, "within": within, "levelled_pairs": len(pairs),
        "mae": round(sum(abs(a - b) for a, b in pairs) / max(len(pairs), 1), 2),
        "per_level": {k: [hit[k], per[k]] for k in LV}, "confusion": {f"{a}>{b}": conf[(a, b)] for a in LV for b in LV},
        "set_by": {k: [src_ok[k], src[k]] for k in src}, "per_concept_fusion_exact": old_exact,
        "concept_matching": {"answers": len(g), "gold_in_linked": in_map, "top_linked_right": top_ok,
                             "judge_main_right": main_ok},
        "linked_counts": link_counts.most_common(), "shown_counts": shown_counts.most_common(),
        "linked_pairs": [[a, b, k] for (a, b), k in pair_counts.most_common(40)],
        "concepts_per_answer": round(sum(len(r["linked"]) for r in rows) / n, 2),
        "rows": rows,
    }
    (run / "report_100.json").write_text(json.dumps(rep, indent=1, ensure_ascii=False), encoding="utf-8")

    print(f"level: exact {exact}/{n} ({exact / n:.0%}), within one {within}/{len(pairs)}, MAE {rep['mae']}, "
          f"per-concept fusion (old final) {old_exact}/{n}")
    print("  per level:", " ".join(f"{k} {hit[k]}/{per[k]}" for k in LV))
    print("  set by:", {k: f"{src_ok[k]}/{src[k]}" for k in src})
    print(f"concepts ({len(g)} hand-checked): right concept linked {in_map}, top-linked right {top_ok}, "
          f"judge's main concept right {main_ok}")

    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "Steps per answer"
    head = ["ID", "Student answer", "Step 1 no-claim check", "Step 2 concepts linked (link score)",
            "Step 2 extra concept (closest rules)", "Step 3 concept map: concepts linked in the graph",
            "Step 4 code-run facts", "Step 5 judge level per concept", "Step 5 judge main concept",
            "Step 5 judge overall level", "Step 5 judge why", "Step 6 final level", "Set by",
            "Classified level", "Match", "Classified why", "Hand-checked concept"]
    a.append(head)
    for r in rows:
        names = [c for c, _ in r["linked"]] + r["extra"]
        edges = [f"{x}-{y}" for x in names for y in names if x < y and (y in kg.concepts[x].neighbours)]
        a.append([r["id"], r["text"], "no claim" if set(r["gates"]) & GATES else "claim",
                  ", ".join(f"{c} ({s})" for c, s in r["linked"]), ", ".join(r["extra"]), ", ".join(edges),
                  "; ".join(r["run_facts"]), ", ".join(f"{c}: {l}" for c, l in r["judge_concepts"].items()),
                  r["judge_main"], r["judge_answer"], r["judge_why"], r["final"], r["set_by"], r["label"],
                  r["match"], r["label_why"], ", ".join(r["gold"])])
    for cell in a[1]:
        cell.font = Font(bold=True)
    for row in a.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHIJKLMNOPQ", (7, 45, 10, 40, 20, 35, 45, 35, 18, 10, 45, 10, 18, 10, 8, 45, 22)):
        a.column_dimensions[col].width = w
    a.freeze_panes = "C2"
    c = wb.create_sheet("Confusion")
    c.append(["classified \\ final"] + LV)
    for x in LV:
        c.append([x] + [conf[(x, y)] for y in LV])
    wb.save(run / "steps_100_rubric.xlsx")
    print("wrote", run / "report_100.json", "and", run / "steps_100_rubric.xlsx")


if __name__ == "__main__":
    main()
